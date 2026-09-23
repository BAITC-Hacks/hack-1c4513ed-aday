"""Normalize the six 1C exports into keyed data frames. Only this module reads files."""
from pathlib import Path
import re
import pickle
import statistics
import hashlib
import pandas as pd
import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FILES = ("sales_tx.xlsx", "sales_monthly.xlsx", "stock_monthly.xlsx", "in_transit.xlsx", "moq.xlsx", "seasonality.xlsx")
MONTHS = ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сент", "окт", "ноя", "дек"]
FULL_MONTHS = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа", "сентября", "октября", "ноября", "декабря"]


def code(x):
    if pd.isna(x):
        return ""
    return str(x).strip().removesuffix(".0")


def number(x, default=0):
    v = pd.to_numeric(x, errors="coerce")
    return default if pd.isna(v) else float(v)


def month_header(x):
    s = str(x).lower().strip()
    year = re.search(r"20\d{2}", s)
    if year:
        for i, stem in enumerate(MONTHS, 1):
            if s.startswith(stem):
                return pd.Timestamp(int(year.group()), i, 1)
    return None


def sheet(path, required):
    """Find a header row by keyword groups, preferring the first matching sheet."""
    with pd.ExcelFile(path) as book:
        for name in book.sheet_names:
            preview = pd.read_excel(book, sheet_name=name, header=None, nrows=8)
            for i, row in preview.iterrows():
                cells = [str(v).lower().strip() for v in row if pd.notna(v)]
                if all(any(key in cell for cell in cells) for key in required):
                    df = pd.read_excel(book, sheet_name=name, header=i)
                    df.columns = [str(c).strip() for c in df.columns]
                    return df
    raise ValueError(f"Не найдены заголовки {required}: {path}")


def column(df, *words):
    for word in words:
        matches = [c for c in df if word in str(c).lower()]
        if matches:
            return matches[0]
    raise KeyError(f"Нет колонки {words}")


def monthly(path):
    df = sheet(path, ["номенклатура", "код"])
    key = column(df, "номенклатура.код", "код 1с")
    name = column(df, "номенклатура")
    months = {c: month_header(c) for c in df if month_header(c) is not None}
    if not months:
        raise ValueError("не найдены колонки месяцев")
    base = df[[key, name, *months]].copy()
    base.columns = ["code", "name", *months.values()]
    base["code"] = base.code.map(code)
    base = base[(base.code != "") & (base.name.astype(str).str.strip().str.lower() != "итого")]
    long = base.melt(id_vars=["code", "name"], var_name="month", value_name="qty")
    long["qty"] = pd.to_numeric(long.qty, errors="coerce").fillna(0)
    return long.groupby(["code", "month"], as_index=False).qty.sum(), base[["code", "name"]].drop_duplicates("code")


def transactions(path):
    df = sheet(path, ["дата", "документ", "количество"])
    cols = {name: column(df, word) for name, word in [("date", "дата"), ("document", "документ"), ("invoice", "номер"), ("code", "код"), ("qty", "количество")]}
    df = df.rename(columns={v: k for k, v in cols.items()})[list(cols)]
    df = df[df.document.astype(str).str.startswith("Расходная накладная")].copy()
    df["date"] = pd.to_datetime(df.date, format="%d.%m.%Y %H:%M:%S", errors="coerce")
    df = df[df.date >= "2025-01-01"]
    df["code"] = df.code.map(code)
    df["invoice"] = df.invoice.map(code)
    df["qty"] = pd.to_numeric(df.qty, errors="coerce").fillna(0)
    return df[["date", "invoice", "code", "qty"]].reset_index(drop=True)


def brand_season(path):
    book = load_workbook(path, read_only=True, data_only=True)
    rows = list(book.active.values)
    years = []
    for row in rows:
        if row and str(row[0]) in ("2024", "2025"):
            years.append([number(x) for x in row[1:13]])
    book.close()
    if not years:
        raise ValueError(f"нет строк сезонности 2024–2025: {path}")
    arr = np.array(years, dtype=float).mean(axis=0)
    arr = arr / arr.mean() if arr.mean() > 0 else np.ones(12)
    return pd.DataFrame({"month_num": range(1, 13), "index": arr})


def transit(path, supplier):
    df = sheet(path, ["код 1с", "наименование"])
    df["code"] = df[column(df, "код 1с")].map(code)
    if supplier == "IEK":
        doc_cols = [c for c in df if "поступление до" in c.lower()]
        if not doc_cols:
            raise ValueError("нет колонок документов с датами поступления")
        lead = []
        for c in doc_cols:
            start = re.search(r"от\s+(\d{1,2})\s+([а-я]+)\s+(20\d{2})", c.lower())
            end = re.search(r"поступление до\s+(\d{2}\.\d{2}\.20\d{2})", c.lower())
            if start and end and start.group(2) in FULL_MONTHS:
                d = pd.Timestamp(int(start.group(3)), FULL_MONTHS.index(start.group(2)) + 1, int(start.group(1)))
                lead.append((pd.to_datetime(end.group(1), dayfirst=True) - d).days)
        df["transit"] = df[doc_cols].apply(pd.to_numeric, errors="coerce").fillna(0).sum(axis=1)
        df["article"] = df[column(df, "артикул")].fillna("").astype(str)
        df["manager_category"] = ""
        df["manager_order"] = 0.0
        df["free_stock"] = np.nan
        lead_days = round(statistics.median(lead)) if lead else 30
    else:
        df["transit"] = pd.to_numeric(df[column(df, "сэ в пути")], errors="coerce").fillna(0)
        df["article"] = df[column(df, "артикул поставщика")].fillna("").astype(str)
        df["manager_category"] = df[column(df, "категория 2026")].fillna("").astype(str)
        df["manager_order"] = pd.to_numeric(df[column(df, "заказ")], errors="coerce").fillna(0)
        df["free_stock"] = pd.to_numeric(df[column(df, "свободный остаток")], errors="coerce")
        lead_days = 30
    return df[["code", "article", "transit", "manager_category", "manager_order", "free_stock"]].drop_duplicates("code"), lead_days


def moq(path):
    # Explicitly read cached formula results, never external VLOOKUP references.
    book = load_workbook(path, read_only=True, data_only=True)
    rows = list(book.active.values)
    book.close()
    header = next(i for i, row in enumerate(rows[:8]) if any("код" in str(v).lower() for v in row) and any("артикул" in str(v).lower() for v in row))
    df = pd.DataFrame(rows[header + 1:], columns=[str(v).strip() for v in rows[header]])
    key = column(df, "код 1с", "номенклатура.код")
    value = column(df, "мин. разр", "кратность")
    df["code"] = df[key].map(code)
    df["moq"] = pd.to_numeric(df[value], errors="coerce").fillna(1).clip(lower=1)
    df["moq"] = np.ceil(df.moq)
    df["moq_article"] = df[column(df, "артикул")].fillna("").astype(str)
    return df[["code", "moq", "moq_article"]].drop_duplicates("code")


def load_supplier(supplier, use_cache=True, source_dir=None):
    if supplier not in ("IEK", "SystemeElectric"):
        raise ValueError(supplier)
    raw = Path(source_dir).resolve() if source_dir is not None else ROOT / "data" / "raw" / supplier
    suffix = "" if source_dir is None else "-" + hashlib.sha256(str(raw).encode()).hexdigest()[:12]
    cache = ROOT / "data" / "cache" / f"{supplier}{suffix}.pkl"
    sources = [raw / name for name in FILES]
    missing = [path.name for path in sources if not path.is_file()]
    if missing:
        raise ValueError(f"Не найдены файлы: {', '.join(missing)}")
    if use_cache and cache.exists() and all(cache.stat().st_mtime > p.stat().st_mtime for p in sources):
        with cache.open("rb") as f:
            return pickle.load(f)
    try:
        sales, names = monthly(raw / "sales_monthly.xlsx")
    except Exception as exc:
        raise ValueError(f"sales_monthly.xlsx: {exc}") from exc
    try:
        stocks, stock_names = monthly(raw / "stock_monthly.xlsx")
    except Exception as exc:
        raise ValueError(f"stock_monthly.xlsx: {exc}") from exc
    try:
        tx = transactions(raw / "sales_tx.xlsx")
    except Exception as exc:
        raise ValueError(f"sales_tx.xlsx: {exc}") from exc
    try:
        transit_df, lead = transit(raw / "in_transit.xlsx", supplier)
    except Exception as exc:
        raise ValueError(f"in_transit.xlsx: {exc}") from exc
    try:
        moq_df = moq(raw / "moq.xlsx")
    except Exception as exc:
        raise ValueError(f"moq.xlsx: {exc}") from exc
    items = pd.concat([names, stock_names], ignore_index=True).drop_duplicates("code")
    items = items.merge(transit_df, on="code", how="left").merge(moq_df, on="code", how="left")
    items["article"] = items.article.fillna("")
    items.loc[items.article == "", "article"] = items.moq_article.fillna("")
    items["moq"] = items.moq.fillna(1)
    for col in ("transit", "manager_order"):
        items[col] = items[col].fillna(0)
    items["manager_category"] = items.manager_category.fillna("")
    items["supplier"] = "IEK" if supplier == "IEK" else "Systeme Electric"
    try:
        season = brand_season(raw / "seasonality.xlsx")
    except Exception as exc:
        raise ValueError(f"seasonality.xlsx: {exc}") from exc
    result = {"items": items.drop(columns="moq_article"), "sales": sales, "stocks": stocks.rename(columns={"qty": "stock"}), "tx": tx, "season": season, "lead_days": lead}
    if use_cache:
        cache.parent.mkdir(parents=True, exist_ok=True)
        with cache.open("wb") as f:
            pickle.dump(result, f)
    return result


def source_signature(supplier, source_dir=None):
    """Cache key changes whenever a source workbook changes."""
    folder = Path(source_dir).resolve() if source_dir is not None else ROOT / "data" / "raw" / supplier
    return (str(folder),) + tuple((name, (folder / name).stat().st_mtime_ns, (folder / name).stat().st_size) for name in FILES)
