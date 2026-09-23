"""Build the approved purchasing workbook without changing order calculations."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ORDER_COLUMNS = ["Код 1С", "Артикул поставщика", "Наименование", "Количество", "Срочность", "Обоснование"]


def build_workbook(approved, as_of):
    book = Workbook()
    summary = book.active
    summary.title = "Сводка"
    summary.append(["Поставщик", "Число позиций", "Дата расчёта"])
    for supplier, group in approved.groupby("Поставщик", sort=True):
        summary.append([supplier, len(group), as_of.strftime("%d.%m.%Y")])
        sheet = book.create_sheet(supplier)
        sheet.append(ORDER_COLUMNS)
        for values in group[ORDER_COLUMNS].itertuples(index=False, name=None):
            sheet.append(list(values))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions["A"].width = 18
        sheet.column_dimensions["B"].width = 25
        sheet.column_dimensions["C"].width = 65
        sheet.column_dimensions["D"].width = 16
        sheet.column_dimensions["E"].width = 18
        sheet.column_dimensions["F"].width = 95
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height = 48
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    for sheet in book:
        sheet.row_dimensions[1].height = 28
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="193B5A")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        if sheet == summary:
            for col, width in enumerate((28, 20, 22), 1):
                sheet.column_dimensions[get_column_letter(col)].width = width
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
