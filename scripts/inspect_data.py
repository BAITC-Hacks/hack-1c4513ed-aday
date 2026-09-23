"""Print sheet names, dimensions and first six rows of every source workbook."""
from pathlib import Path
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1] / "data" / "raw"


def inspect() -> None:
    for path in sorted(ROOT.glob("*/*.xlsx")):
        print(f"\n{path.relative_to(ROOT)} ({path.stat().st_size:,} bytes)")
        book = load_workbook(path, read_only=True, data_only=True)
        for sheet in book:
            print(f"  {sheet.title}: {sheet.max_row} rows x {sheet.max_column} columns")
            for row in sheet.iter_rows(max_row=6, values_only=True):
                print("   ", repr(row))
        book.close()


if __name__ == "__main__":
    inspect()
