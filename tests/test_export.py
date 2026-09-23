from datetime import date
from io import BytesIO
import pandas as pd
from openpyxl import load_workbook
from src.export import build_workbook, ORDER_COLUMNS


def test_approved_workbook_has_supplier_sheets_and_summary():
    rows = pd.DataFrame({
        "Поставщик": ["IEK", "Systeme Electric"],
        "Код 1С": ["001_", "002_"],
        "Артикул поставщика": ["A", "B"],
        "Наименование": ["Товар А", "Товар Б"],
        "Количество": [12, 24],
        "Срочность": ["Критично", "Плановая"],
        "Обоснование": ["Обоснование А", "Обоснование Б"],
    })
    book = load_workbook(BytesIO(build_workbook(rows, date(2026, 9, 22))))
    assert book.sheetnames == ["Сводка", "IEK", "Systeme Electric"]
    assert [c.value for c in book["IEK"][1]] == ORDER_COLUMNS
    assert book["IEK"]["D2"].value == 12
    assert book["Systeme Electric"]["F2"].value == "Обоснование Б"
    assert book["Сводка"]["B2"].value == 1
    assert book["Сводка"]["C2"].value == "22.09.2026"
    assert book["IEK"]["A1"].font.bold
    assert book["IEK"].column_dimensions["F"].width >= 80
